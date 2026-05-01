"""Build EXTENDED climate dataset for ~1065 cities.

Inputs:
    outputs/lsk-lskos-cities.json    — registry (city, region, snow_region, sg_kpa)
    outputs/ese-data-extended.json   — scraped ese.pro values
    outputs/geocoded-cities.json     — Nominatim lat/lon

Outputs:
    src/data/regions/settlements-climate-extended.json
    data-source/settlements-climate-extended-master.xlsx
"""

from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
LSK_PATH = ROOT / "outputs" / "lsk-lskos-cities.json"
ESE_PATH = ROOT / "outputs" / "ese-data-extended.json"
GEO_PATH = ROOT / "outputs" / "geocoded-cities.json"

JSON_OUT = ROOT / "src" / "data" / "regions" / "settlements-climate-extended.json"
XLSX_OUT = ROOT / "data-source" / "settlements-climate-extended-master.xlsx"

SOURCE_LSK = "lsk-lskos.ru/sr (СП 20.13330.2016, Изм. №2, табл. 10.1)"
SOURCE_ESE = (
    "СП 20.13330.2016 / СП 14.13330 / СП 131.13330 / СП 22.13330; "
    "сверено через ese.pro/tools/calculatory/klimaticheskie-nagruzki/"
)
SOURCE_OSM = "OSM Nominatim"

ALLOWED_TERRAIN = ["A", "B", "C"]

_TRANS = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e",
    "ж": "zh", "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m",
    "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u",
    "ф": "f", "х": "kh", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "shch",
    "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
}


def slugify(s: str) -> str:
    out = []
    for ch in s.strip().lower():
        if ch in _TRANS:
            out.append(_TRANS[ch])
        elif ch.isalnum():
            out.append(ch)
        elif ch in (" ", "-", "_", "."):
            out.append("_")
    res = "".join(out)
    return re.sub(r"_+", "_", res).strip("_")


def make_id(city: str, region: str) -> str:
    return f"{slugify(city)}__{slugify(region)}"


def make_geo_key(city: str, region: str) -> str:
    return f"{city}||{region}"


def _round_sg(kpa: float | None) -> float | None:
    """Snap exact Sg to its standard district value."""
    if kpa is None:
        return None
    table = [
        (0.5, "I"), (1.0, "II"), (1.5, "III"), (2.0, "IV"),
        (2.5, "V"), (3.0, "VI"), (3.5, "VII"), (4.0, "VIII"),
    ]
    best = min(table, key=lambda kv: abs(kv[0] - kpa))
    return best[0]


def build_records() -> list[dict[str, Any]]:
    cities = json.load(LSK_PATH.open("r", encoding="utf-8"))
    ese = json.load(ESE_PATH.open("r", encoding="utf-8")) if ESE_PATH.exists() else {}
    geo = json.load(GEO_PATH.open("r", encoding="utf-8")) if GEO_PATH.exists() else {}

    records: list[dict[str, Any]] = []

    for c in cities:
        city = c["city"]
        region = c["region"]
        sid = make_id(city, region)
        ese_row = ese.get(sid) or {}
        ese_data = (ese_row or {}).get("data") or {}
        geo_row = geo.get(make_geo_key(city, region)) or {}

        # ── snow ─────────────────────────────────────────────────────────
        # primary source: ese.pro (district + Sg). Fallback: lsk-lskos (Sg only)
        ese_snow = ese_data.get("loads", {}).get("snow", {}) or {}
        ese_snow_kpa = ese_snow.get("kpa")
        ese_snow_region = ese_snow.get("region")

        if ese_snow_region and ese_snow_kpa is not None:
            sg_district = _round_sg(ese_snow_kpa)
            snow_block = {
                "region": ese_snow_region,
                "sgKpa": sg_district,
                "sgKpaExact": ese_snow_kpa,
                "source": SOURCE_ESE,
                "status": "verified",
            }
        elif c.get("sg_kpa") is not None:
            snow_block = {
                "region": c.get("snow_region"),
                "sgKpa": c.get("sg_kpa"),
                "sgKpaExact": None,
                "source": SOURCE_LSK,
                "status": "from_sp20_appendix_k",
            }
        else:
            snow_block = {
                "region": None,
                "sgKpa": None,
                "sgKpaExact": None,
                "source": SOURCE_LSK,
                "status": "requires_verification",
            }

        # cross-check: does ese match lsk?
        snow_match = None
        if (
            c.get("sg_kpa") is not None
            and ese_snow_region is not None
            and snow_block["sgKpa"] is not None
        ):
            snow_match = (
                c.get("snow_region") == ese_snow_region
                and abs(c["sg_kpa"] - snow_block["sgKpa"]) < 0.01
            )

        # ── wind ─────────────────────────────────────────────────────────
        wind = ese_data.get("loads", {}).get("wind", {}) or {}
        if wind.get("region") and wind.get("kpa") is not None:
            wind_block = {
                "region": wind["region"],
                "w0Kpa": wind["kpa"],
                "source": SOURCE_ESE,
                "status": "verified",
            }
        else:
            wind_block = {
                "region": None,
                "w0Kpa": None,
                "source": SOURCE_ESE,
                "status": "requires_verification",
            }

        # ── ice ──────────────────────────────────────────────────────────
        ice = ese_data.get("loads", {}).get("ice", {}) or {}
        if ice.get("region") and ice.get("mm") is not None:
            ice_block = {
                "region": ice["region"],
                "iceThicknessMm": ice["mm"],
                "source": SOURCE_ESE,
                "status": "verified",
            }
        else:
            ice_block = {
                "region": None,
                "iceThicknessMm": None,
                "source": SOURCE_ESE,
                "status": "requires_verification",
            }

        # ── seismic ─────────────────────────────────────────────────────
        seismic = ese_data.get("seismic", {}) or {}
        if seismic.get("mapA") is not None:
            seismic_block = {
                "points": seismic["mapA"],
                "pointsMapB": seismic.get("mapB"),
                "pointsMapC": seismic.get("mapC"),
                "source": SOURCE_ESE,
                "status": "verified",
            }
        else:
            seismic_block = {
                "points": None,
                "pointsMapB": None,
                "pointsMapC": None,
                "source": SOURCE_ESE,
                "status": "requires_verification",
            }

        # ── coordinates ─────────────────────────────────────────────────
        if geo_row.get("lat") is not None and geo_row.get("lon") is not None:
            coords_block = {
                "lat": geo_row["lat"],
                "lon": geo_row["lon"],
                "displayName": geo_row.get("display_name"),
                "source": SOURCE_OSM,
                "status": "verified",
            }
        else:
            coords_block = {
                "lat": None,
                "lon": None,
                "displayName": None,
                "source": SOURCE_OSM,
                "status": geo_row.get("status") or "requires_verification",
            }

        # ── frost depth ─────────────────────────────────────────────────
        fd = ese_data.get("frostDepth", {}) or {}
        if any(fd.get(k) is not None for k in ("loamy", "sandyFine", "sandyCoarse", "coarseFragmental")):
            frost_block = {
                "loamy": fd.get("loamy"),
                "sandyFine": fd.get("sandyFine"),
                "sandyCoarse": fd.get("sandyCoarse"),
                "coarseFragmental": fd.get("coarseFragmental"),
                "source": SOURCE_ESE,
                "status": "verified",
            }
        else:
            frost_block = None

        # ── cold period ─────────────────────────────────────────────────
        cold = ese_data.get("cold", {}) or {}
        cold_block = {**cold, "source": SOURCE_ESE, "status": "verified"} if cold else None

        # ── warm period ─────────────────────────────────────────────────
        warm = ese_data.get("warm", {}) or {}
        warm_block = {**warm, "source": SOURCE_ESE, "status": "verified"} if warm else None

        # ── monthly temps ───────────────────────────────────────────────
        mt = ese_data.get("monthlyTemps")
        if mt and len(mt) == 13:
            monthly_block = {
                "byMonth": [float(x) if x is not None else None for x in mt[:12]],
                "yearly": float(mt[12]) if mt[12] is not None else None,
                "source": SOURCE_ESE,
                "status": "verified",
            }
        else:
            monthly_block = None

        # ── data status & comment ───────────────────────────────────────
        statuses = [snow_block["status"], wind_block["status"],
                    ice_block["status"], seismic_block["status"]]
        if all(s == "verified" for s in statuses):
            data_status = "verified"
        elif any(s == "verified" for s in statuses):
            data_status = "partial"
        else:
            data_status = "requires_verification"

        comment_parts: list[str] = []
        if snow_match is False:
            comment_parts.append(
                f"⚠ Расхождение Sg: lsk-lskos {c.get('snow_region')}/{c.get('sg_kpa')} vs "
                f"ese.pro {ese_snow_region}/{snow_block['sgKpa']}"
            )
        elif snow_match is True:
            comment_parts.append("Sg сверено с lsk-lskos и ese.pro — совпадает")

        if (
            snow_block["sgKpa"] is not None
            and snow_block.get("sgKpaExact") is not None
            and abs(snow_block["sgKpaExact"] - snow_block["sgKpa"]) > 0.01
        ):
            comment_parts.append(
                f"Sg уточнённое ese.pro: {snow_block['sgKpaExact']} кПа"
            )

        rec = {
            "id": sid,
            "country": "Россия",
            "region": region,
            "settlement": city,
            "settlementType": None,
            "coordinates": coords_block,
            "snow": snow_block,
            "wind": wind_block,
            "ice": ice_block,
            "seismic": seismic_block,
            "terrain": {
                "defaultType": None,
                "allowedTypes": ALLOWED_TERRAIN,
            },
            "manualAllowed": True,
            "expertOverrideAllowed": True,
            "dataStatus": data_status,
            "comment": " ".join(comment_parts),
        }
        if frost_block is not None:
            rec["frostDepth"] = frost_block
        if cold_block is not None:
            rec["coldPeriod"] = cold_block
        if warm_block is not None:
            rec["warmPeriod"] = warm_block
        if monthly_block is not None:
            rec["monthlyTemps"] = monthly_block

        records.append(rec)

    records.sort(key=lambda r: (r["region"], r["settlement"]))
    return records


# ── Excel writer ────────────────────────────────────────────────────────────

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
    ws.freeze_panes = ws[f"A{row + 1}"]


def autosize(ws, min_w: int = 10, max_w: int = 40) -> None:
    for i, col in enumerate(ws.columns, start=1):
        length = min_w
        for cell in col:
            v = cell.value
            if v is None:
                continue
            length = max(length, min(max_w, len(str(v)) + 2))
        ws.column_dimensions[get_column_letter(i)].width = length


MASTER_HEADERS = [
    "id", "country", "region", "settlement",
    "lat", "lon",
    "snow_region", "sg_kpa", "sg_kpa_exact", "snow_source", "snow_status",
    "wind_region", "w0_kpa", "wind_source", "wind_status",
    "ice_region", "ice_thickness_mm", "ice_source", "ice_status",
    "seismic_points", "seismic_points_b", "seismic_points_c",
    "seismic_source", "seismic_status",
    "terrain_type_default", "terrain_type_allowed",
    "manual_allowed", "expert_override_allowed",
    "data_status", "comment",
]

EXTRAS_HEADERS = [
    "id", "settlement", "region",
    "frost_loamy", "frost_sandy_fine", "frost_sandy_coarse", "frost_coarse_fragm",
    "t_coldest_day_098", "t_coldest_day_092",
    "t_coldest_5days_098", "t_coldest_5days_092",
    "t_094", "t_abs_min", "amp_cold_month",
    "dur_le0", "mean_le0", "dur_le8", "mean_le8", "dur_le10", "mean_le10",
    "humidity_cold", "humidity_cold_15",
    "precip_nov_mar", "wind_dir_dec_feb", "max_wind_jan", "mean_wind_le8",
    "barometric_hpa", "t_095", "t_099",
    "mean_max_t_warm", "t_abs_max", "amp_warm_month",
    "humidity_warm", "humidity_warm_15",
    "precip_apr_oct", "daily_max_precip",
    "wind_dir_jun_aug", "min_wind_jul",
    "t_jan", "t_feb", "t_mar", "t_apr", "t_may", "t_jun",
    "t_jul", "t_aug", "t_sep", "t_oct", "t_nov", "t_dec", "t_year",
]

CROSSCHECK_HEADERS = [
    "id", "settlement", "region",
    "lsk_snow_region", "lsk_sg_kpa",
    "ese_snow_region", "ese_sg_kpa", "ese_sg_kpa_exact",
    "match",
]


def row_master(s: dict[str, Any]) -> list[Any]:
    coords = s.get("coordinates") or {}
    seismic = s["seismic"]
    return [
        s["id"], s["country"], s["region"], s["settlement"],
        coords.get("lat"), coords.get("lon"),
        s["snow"]["region"], s["snow"]["sgKpa"], s["snow"].get("sgKpaExact"),
        s["snow"]["source"], s["snow"]["status"],
        s["wind"]["region"], s["wind"]["w0Kpa"],
        s["wind"]["source"], s["wind"]["status"],
        s["ice"]["region"], s["ice"]["iceThicknessMm"],
        s["ice"]["source"], s["ice"]["status"],
        seismic["points"], seismic.get("pointsMapB"), seismic.get("pointsMapC"),
        seismic["source"], seismic["status"],
        s["terrain"]["defaultType"],
        ", ".join(s["terrain"]["allowedTypes"]),
        s["manualAllowed"], s["expertOverrideAllowed"],
        s["dataStatus"], s["comment"],
    ]


def row_extras(s: dict[str, Any]) -> list[Any]:
    fd = s.get("frostDepth") or {}
    cp = s.get("coldPeriod") or {}
    wp = s.get("warmPeriod") or {}
    mt = s.get("monthlyTemps") or {}
    by_month = (mt.get("byMonth") or [None] * 12) if mt else [None] * 12
    return [
        s["id"], s["settlement"], s["region"],
        fd.get("loamy"), fd.get("sandyFine"), fd.get("sandyCoarse"), fd.get("coarseFragmental"),
        cp.get("tempColdestDay098"), cp.get("tempColdestDay092"),
        cp.get("tempColdest5days098"), cp.get("tempColdest5days092"),
        cp.get("temp094"), cp.get("absMin"), cp.get("dailyAmplitude"),
        cp.get("durationLe0"), cp.get("meanLe0"),
        cp.get("durationLe8"), cp.get("meanLe8"),
        cp.get("durationLe10"), cp.get("meanLe10"),
        cp.get("humidityCold"), cp.get("humidityCold15"),
        cp.get("precipNovMar"),
        cp.get("prevailingWindDecFeb"),
        cp.get("maxWindJan"), cp.get("meanWindLe8"),
        wp.get("barometric"),
        wp.get("temp095"), wp.get("temp099"),
        wp.get("meanMaxTempWarmMonth"), wp.get("absMax"),
        wp.get("dailyAmplitude"),
        wp.get("humidityWarm"), wp.get("humidityWarm15"),
        wp.get("precipAprOct"), wp.get("dailyMaxPrecip"),
        wp.get("prevailingWindJunAug"), wp.get("minWindJul"),
        *by_month, mt.get("yearly"),
    ]


def build_crosscheck(records: list[dict[str, Any]]) -> list[list[Any]]:
    cities = json.load(LSK_PATH.open("r", encoding="utf-8"))
    by_id = {make_id(c["city"], c["region"]): c for c in cities}

    rows: list[list[Any]] = []
    for r in records:
        lsk = by_id.get(r["id"])
        if not lsk:
            continue
        snow = r["snow"]
        ese_region = snow["region"] if snow["status"] == "verified" else None
        ese_kpa = snow["sgKpa"] if snow["status"] == "verified" else None
        match: str
        if ese_region and lsk.get("sg_kpa") is not None and ese_kpa is not None:
            ok_region = (lsk.get("snow_region") == ese_region)
            ok_kpa = abs(lsk["sg_kpa"] - ese_kpa) < 0.01
            match = "ok" if (ok_region and ok_kpa) else "DIFF"
        else:
            match = "no_ese_data"
        rows.append([
            r["id"], r["settlement"], r["region"],
            lsk.get("snow_region"), lsk.get("sg_kpa"),
            ese_region, ese_kpa, snow.get("sgKpaExact"),
            match,
        ])
    return rows


def write_xlsx(records: list[dict[str, Any]]) -> None:
    XLSX_OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "01_master_settlements"
    ws.append(MASTER_HEADERS)
    for r in records:
        ws.append(row_master(r))
    style_header(ws)
    autosize(ws)

    ws_extras = wb.create_sheet("05_climate_extras")
    ws_extras.append(EXTRAS_HEADERS)
    for r in records:
        ws_extras.append(row_extras(r))
    style_header(ws_extras)
    autosize(ws_extras)

    ws_check = wb.create_sheet("06_snow_crosscheck")
    ws_check.append(CROSSCHECK_HEADERS)
    for row in build_crosscheck(records):
        ws_check.append(row)
    style_header(ws_check)
    autosize(ws_check)

    ws_src = wb.create_sheet("04_sources_plan")
    ws_src.append(["Источник", "Назначение", "Статус"])
    for s in [
        ("lsk-lskos.ru/sr",
         "Реестр городов + Sg по СП 20 (Изм. №2, табл. 10.1)",
         "источник списка городов и сверочный для Sg"),
        ("ese.pro/tools/calculatory/klimaticheskie-nagruzki/",
         "Снег / ветер / гололёд / сейсмика / СП 22 / СП 131 / среднемесячные T",
         "основной источник значений"),
        ("OSM Nominatim",
         "Координаты (lat/lon) для отрисовки на карте",
         "вспомогательный"),
        ("СП 20.13330.2016", "Карты районирования + табл. 10.1", "первоисточник"),
        ("СП 14.13330", "Сейсмика (ОСР-2015 карты A/B/C)", "первоисточник"),
        ("СП 131.13330", "Климатические параметры периодов", "первоисточник"),
        ("СП 22.13330", "Глубина сезонного промерзания", "первоисточник"),
    ]:
        ws_src.append(list(s))
    style_header(ws_src)
    autosize(ws_src)

    wb.save(XLSX_OUT)


def main() -> int:
    records = build_records()
    JSON_OUT.parent.mkdir(parents=True, exist_ok=True)
    JSON_OUT.write_text(
        json.dumps(records, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    write_xlsx(records)

    n_total = len(records)
    n_verified = sum(1 for r in records if r["dataStatus"] == "verified")
    n_partial = sum(1 for r in records if r["dataStatus"] == "partial")
    n_pending = sum(1 for r in records if r["dataStatus"] == "requires_verification")
    n_geo = sum(1 for r in records if (r.get("coordinates") or {}).get("lat") is not None)
    n_frost = sum(1 for r in records if r.get("frostDepth"))
    n_cold = sum(1 for r in records if r.get("coldPeriod"))

    print(f"settlements: {n_total}")
    print(f"  verified  : {n_verified}")
    print(f"  partial   : {n_partial}")
    print(f"  pending   : {n_pending}")
    print(f"  geo       : {n_geo}")
    print(f"  frostDepth: {n_frost}")
    print(f"  coldPeriod: {n_cold}")
    print(f"json:  {JSON_OUT}")
    print(f"xlsx:  {XLSX_OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
